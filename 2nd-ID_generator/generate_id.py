#!/usr/bin/env python3
"""
HPAP 2nd ID Generator

Computes the "2nd ID" for a donor from the pancreas cross-clamp time and
logs the entry into HPAP_2nd_ID_Log.xlsx.

2nd ID format: YY + M + DD + HH  (all from the cross-clamp time, local time)
  YY = 2-digit year
  M  = single-letter month code (custom scheme, avoids Jan/Jun/Jul & Mar/May & Apr/Aug collisions)
  DD = 2-digit day
  HH = 2-digit hour, 24-hour clock

Run it with:  python3 generate_id.py
"""

import datetime
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

LOG_FILE = Path(__file__).parent / "HPAP_2nd_ID_Log.xlsx"
SHEET_NAME = "2nd-ID Log"
HEADERS = ["HPAP_ID", "UNOS#", "cross_clamp_time", "time_zone", "disease_status", "notes", "2nd_ID"]

MONTH_CODE = {
    1: "J", 2: "F", 3: "M", 4: "A", 5: "Y", 6: "U",
    7: "L", 8: "G", 9: "S", 10: "O", 11: "N", 12: "D",
}


def parse_crossclamp(text):
    """Parse a cross-clamp date/time string into a datetime.

    Accepts things like: 10/29/2016 15:13, 10/29/16, 15:13, 1/21/24 4:13,
    3/02/24, 22.02, 2/08/23, 0053
    """
    text = text.strip()

    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})[,\.]?\s*(\d{1,2})[:\.](\d{2})$", text)
    if m:
        mo, d, y, h, mi = m.groups()
    else:
        m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})[,\.]?\s*(\d{3,4})$", text)
        if not m:
            raise ValueError(
                f"Couldn't understand the date/time {text!r}. "
                "Try a format like 8/20/2026 14:30"
            )
        mo, d, y, hm = m.groups()
        hm = hm.zfill(4)
        h, mi = hm[:2], hm[2:]

    year = int(y)
    if year < 100:
        year += 2000

    return datetime.datetime(year, int(mo), int(d), int(h), int(mi))


def compute_second_id(dt):
    return f"{dt.year % 100:02d}{MONTH_CODE[dt.month]}{dt.day:02d}{dt.hour:02d}"


def load_or_create_log():
    if LOG_FILE.exists():
        wb = openpyxl.load_workbook(LOG_FILE)
        ws = wb[SHEET_NAME]
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        for col, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
        for col in range(1, len(HEADERS) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
    return wb, ws


def existing_hpap_ids(ws):
    ids = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            ids.add(str(row[0]).strip())
    return ids


def suggest_next_hpap_id(ws):
    max_num = 0
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        hpap_id = row[0]
        if hpap_id:
            m = re.match(r"HPAP0*(\d+)", str(hpap_id).strip(), re.IGNORECASE)
            if m:
                max_num = max(max_num, int(m.group(1)))
    return f"HPAP{max_num + 1:03d}" if max_num else "HPAP001"


def append_entry(ws, hpap_id, unos, crossclamp_dt, timezone, disease_status, notes, second_id):
    row = [hpap_id, unos, crossclamp_dt, timezone, disease_status, notes, second_id]
    ws.append(row)
    ws.cell(row=ws.max_row, column=3).number_format = "m/d/yy h:mm"


def prompt(label, default=None, required=True):
    suffix = f" [{default}]" if default else ""
    while True:
        val = input(f"{label}{suffix}: ").strip()
        if not val and default is not None:
            val = default
        if val or not required:
            return val
        print("  This field is required.")


def main():
    print("=" * 50)
    print("HPAP 2nd ID Generator")
    print("=" * 50)

    wb, ws = load_or_create_log()
    ids_seen = existing_hpap_ids(ws)

    while True:
        default_id = suggest_next_hpap_id(ws)
        while True:
            hpap_id = prompt("HPAP_ID", default=default_id)
            if hpap_id in ids_seen:
                confirm = input(
                    f"  '{hpap_id}' is already in the log. Add it again anyway? (y/N): "
                ).strip().lower()
                if confirm != "y":
                    continue
            break

        unos = prompt("UNOS#", required=False)

        while True:
            cc_raw = prompt("Cross clamp time, local (e.g. 8/20/2026 14:30)")
            try:
                cc_dt = parse_crossclamp(cc_raw)
                break
            except ValueError as e:
                print(f"  {e}")

        timezone = prompt("Time zone (e.g. EST)", required=False)
        disease_status = prompt("Disease status", required=False)
        notes = prompt("Notes", required=False)

        second_id = compute_second_id(cc_dt)
        append_entry(ws, hpap_id, unos, cc_dt, timezone, disease_status, notes, second_id)
        wb.save(LOG_FILE)
        ids_seen.add(hpap_id)

        print(f"\n  2nd ID: {second_id}")
        print(f"  Saved to {LOG_FILE.name}\n")

        again = input("Add another entry? (Y/n): ").strip().lower()
        if again == "n":
            break

    print("Done.")


if __name__ == "__main__":
    main()
